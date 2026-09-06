import sys
import io
import json
import os
import base64
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm


def _decode_chart_image(base64_str: str) -> io.BytesIO:
    """Decode a base64 PNG string (with or without data URI prefix) into a BytesIO buffer."""
    if "," in base64_str:
        base64_str = base64_str.split(",", 1)[1]
    img_bytes = base64.b64decode(base64_str)
    buf = io.BytesIO(img_bytes)
    buf.seek(0)
    return buf


def generate_pdf(title, subtitle, headers, data, output_path, charts=None):
    num_cols = len(headers)
    pagesize = landscape(A4) if num_cols > 8 else A4

    doc = SimpleDocTemplate(
        output_path,
        pagesize=pagesize,
        rightMargin=1*cm, leftMargin=1*cm, topMargin=1.5*cm, bottomMargin=1.5*cm
    )

    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=22,
        spaceAfter=5,
        textColor=colors.HexColor('#0F172A'),
        alignment=0
    )

    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Normal'],
        fontSize=12,
        spaceAfter=25,
        textColor=colors.HexColor('#64748B'),
        alignment=0
    )

    chart_title_style = ParagraphStyle(
        'ChartTitle',
        parent=styles['Heading2'],
        fontSize=14,
        spaceBefore=10,
        spaceAfter=8,
        textColor=colors.HexColor('#1E293B'),
        alignment=1
    )

    # ── Header section ──
    elements.append(Paragraph(title.upper(), title_style))
    elements.append(Paragraph(subtitle, subtitle_style))
    elements.append(Spacer(1, 15))

    # ── Charts section ──
    if charts:
        avail_width = pagesize[0] - 2 * cm
        for chart_item in charts:
            chart_label = chart_item.get("title", "")
            chart_b64 = chart_item.get("image", "")
            if not chart_b64:
                continue

            elements.append(Paragraph(chart_label, chart_title_style))

            buf = _decode_chart_image(chart_b64)
            img = Image(buf)
            img_width = min(avail_width, 18 * cm)
            aspect = img.imageHeight / img.imageWidth if img.imageWidth else 1
            img_height = img_width * aspect
            img.drawWidth = img_width
            img.drawHeight = img_height
            img.hAlign = 'CENTER'
            elements.append(img)
            elements.append(Spacer(1, 20))

        elements.append(PageBreak())

    # ── Table section ──
    if headers and data:
        header_keys = list(headers.keys())
        header_labels = [headers[k] for k in header_keys]

        table_data = [header_labels]
        for row in data:
            row_values = []
            for k in header_keys:
                val = row.get(k, 0) if k != 'indicador' else row.get(k, '')
                if k == 'indicador' and len(str(val)) > 50:
                    val = str(val)[:47] + '...'
                row_values.append(str(val))
            table_data.append(row_values)

        avail_width = pagesize[0] - 2 * cm
        if num_cols > 1:
            first_col_width = avail_width * 0.45 if pagesize == A4 else avail_width * 0.35
            other_cols_width = (avail_width - first_col_width) / (num_cols - 1)
            col_widths = [first_col_width] + [other_cols_width] * (num_cols - 1)
        else:
            col_widths = [avail_width]

        t = Table(table_data, colWidths=col_widths, repeatRows=1)

        style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ALIGN', (0, 0), (0, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8 if num_cols > 10 else 9),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]

        for i, row in enumerate(data):
            if " > " in str(row.get('indicador', '')):
                style.append(('TEXTCOLOR', (0, i+1), (0, i+1), colors.HexColor('#475569')))
                style.append(('LEFTPADDING', (0, i+1), (0, i+1), 20))
            else:
                style.append(('FONTNAME', (0, i+1), (-1, i+1), 'Helvetica-Bold'))
                style.append(('BACKGROUND', (0, i+1), (-1, i+1), colors.HexColor('#F1F5F9')))

        if 'total' in header_keys:
            total_idx = header_keys.index('total')
            style.append(('FONTNAME', (total_idx, 0), (total_idx, -1), 'Helvetica-Bold'))
            style.append(('BACKGROUND', (total_idx, 1), (total_idx, -1), colors.HexColor('#F1F5F9')))
            style.append(('TEXTCOLOR', (total_idx, 1), (total_idx, -1), colors.HexColor('#1E293B')))

        t.setStyle(TableStyle(style))
        elements.append(t)

    # ── Footer ──
    elements.append(Spacer(1, 40))
    footer_text = f"RELATÓRIO GERADO EM {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - INDICADORES HEALTHMAIS"
    elements.append(Paragraph(
        footer_text,
        ParagraphStyle('Footer', parent=styles['Normal'], fontSize=7, textColor=colors.grey, alignment=1, letterSpacing=1)
    ))

    doc.build(elements)
    return output_path


def generate_pptx(title, subtitle, headers, data, output_path, charts=None):
    """Generate a PowerPoint presentation with charts and data table."""
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    brand_dark = RGBColor(0x1E, 0x29, 0x3B)
    brand_accent = RGBColor(0x3B, 0x82, 0xF6)
    brand_light = RGBColor(0xF8, 0xFA, 0xFC)

    # ── Slide 1: Title ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = brand_dark

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(1))
    tf2 = sub_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(18)
    p2.font.color.rgb = brand_accent
    p2.alignment = PP_ALIGN.CENTER

    # ── Chart slides ──
    if charts:
        for chart_item in charts:
            chart_label = chart_item.get("title", "")
            chart_b64 = chart_item.get("image", "")
            if not chart_b64:
                continue

            slide = prs.slides.add_slide(prs.slide_layouts[6])

            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = chart_label
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = brand_dark
            p.alignment = PP_ALIGN.CENTER

            buf = _decode_chart_image(chart_b64)
            slide.shapes.add_picture(
                buf,
                Inches(0.5), Inches(1.2),
                Inches(12.333), Inches(5.8)
            )

    # ── Table slide ──
    if headers and data:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = "TABELA DE INDICADORES"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = brand_dark
        p.alignment = PP_ALIGN.CENTER

        header_keys = list(headers.keys())
        header_labels = [headers[k] for k in header_keys]
        num_cols = len(header_keys)
        num_rows = min(len(data) + 1, 25)  # limit rows to fit slide

        table_shape = slide.shapes.add_table(
            num_rows, num_cols,
            Inches(0.3), Inches(1.0),
            Inches(12.7), Inches(6.0)
        )
        table = table_shape.table

        # Header row
        for j, label in enumerate(header_labels):
            cell = table.cell(0, j)
            cell.text = label
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.bold = True
                paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                paragraph.alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = brand_dark

        # Data rows
        for i, row in enumerate(data[:num_rows - 1]):
            for j, k in enumerate(header_keys):
                cell = table.cell(i + 1, j)
                val = row.get(k, 0) if k != 'indicador' else row.get(k, '')
                cell.text = str(val)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(8)
                    paragraph.alignment = PP_ALIGN.CENTER if k != 'indicador' else PP_ALIGN.LEFT

                is_parent = " > " not in str(row.get('indicador', ''))
                if is_parent and k != 'indicador':
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = brand_light

    prs.save(output_path)
    return output_path


def main():
    try:
        if hasattr(sys.stdin, 'buffer'):
            sys.stdin = __import__('io').TextIOWrapper(sys.stdin.buffer, encoding='utf-8')

        input_data = sys.stdin.read()
        if not input_data.strip():
            print(json.dumps({"error": "No input data provided"}))
            sys.exit(1)

        payload = json.loads(input_data)

        fmt = payload.get("format", "pdf")
        title = payload.get("title", "RELATÓRIO DE INDICADORES")
        subtitle = payload.get("subtitle", "")
        headers = payload.get("headers", {})
        data = payload.get("data", [])
        charts = payload.get("charts", [])

        temp_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'tmp')
        os.makedirs(temp_dir, exist_ok=True)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        if fmt == "pptx":
            output_filename = f"report_{timestamp}.pptx"
            output_path = os.path.join(temp_dir, output_filename)
            generate_pptx(title, subtitle, headers, data, output_path, charts=charts)
        else:
            output_filename = f"report_{timestamp}.pdf"
            output_path = os.path.join(temp_dir, output_filename)
            generate_pdf(title, subtitle, headers, data, output_path, charts=charts)

        print(json.dumps({"success": True, "filePath": output_path, "format": fmt}))

    except Exception as e:
        import traceback
        print(json.dumps({"error": str(e), "trace": traceback.format_exc()}))
        sys.exit(1)


if __name__ == "__main__":
    main()
